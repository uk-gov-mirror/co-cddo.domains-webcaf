from io import BytesIO
from pathlib import Path

from django.test import TestCase
from openpyxl import load_workbook

from webcaf.webcaf.models import Assessment
from webcaf.webcaf.utils.to_spreadsheet import create_assessment_workbook

FIXTURE_PATH = Path(__file__).parent.parent / "fixtures" / "completed_assessment_base.json"


class TestAssessmentToExcel(TestCase):
    @classmethod
    def setUpTestData(cls):
        from tests.test_utils import helpers

        cls.org, systems, cls.user = helpers.create_org_systems_and_user(num_systems=1)
        cls.system = systems[0]
        base_assessment = helpers.load_fixture_json("completed_assessment_base.json")

        # Create independent review assessment
        cls.assessment = Assessment.objects.create(
            system=cls.system,
            status="submitted",
            assessment_period="25/26",
            review_type="independent",
            framework="caf32",
            caf_profile="baseline",
            assessments_data=base_assessment,
        )

    def test_create_assessment_workbook_generates_expected_sheets_and_metadata(self):
        wb = create_assessment_workbook(self.assessment)

        # Save to bytes and reload to validate file integrity
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        loaded = load_workbook(out)

        # Expected sheets
        self.assertIn("Self-assessment details", loaded.sheetnames)
        self.assertIn("IGPs", loaded.sheetnames)
        self.assertIn("Contributing outcomes", loaded.sheetnames)

        # Check metadata values
        ws_meta = loaded["Self-assessment details"]
        self.assertEqual(ws_meta["B1"].value, "Test Organisation")
        self.assertEqual(ws_meta["B2"].value, "Test System")
        self.assertEqual(ws_meta["B3"].value, "25/26")

    def test_igp_tab_contains_indicator_answers(self):
        wb = create_assessment_workbook(self.assessment)
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        loaded = load_workbook(out)

        ws_igp = loaded["IGPs"]
        # Header should contain Self-assessment column
        headers = [cell.value for cell in ws_igp[1]]
        self.assertEqual(
            ["Contributing outcome", "IGP", "IGP wording", "Self-assessment", "Self-assessment comments"], headers
        )

        # There should be a data row for the single indicator with 'Y' for True
        # Find first data row
        row2 = [cell.value for cell in ws_igp[2]]
        # Self-assessment column is the 4th column
        self.assertEqual(
            row2,
            [
                "A1.a Board Direction",
                "A1.a Achieved statement 1",
                "Your organisation's approach and policy relating to the security of network and information systems supporting the operation of your essential function(s) are owned and managed at board-level. These are communicated, in a meaningful way, to risk management decision-makers across the organisation.",
                "Y",
                None,
            ],
        )

    def test_contributing_outcomes_has_profile_and_summary(self):
        wb = create_assessment_workbook(self.assessment)
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        loaded = load_workbook(out)

        ws_co = loaded["Contributing outcomes"]
        # Header row should include 'Contributing outcome' and 'Target CAF profile'
        headers = [cell.value for cell in ws_co[1]]
        self.assertIn("Contributing outcome", headers)
        self.assertIn("Target CAF profile", headers)

        # Check first data row contains the contributing outcome title and profile met value
        row2 = [cell.value for cell in ws_co[2]]
        self.assertEqual(
            row2,
            [
                "A1.a Board Direction",
                "Achieved",
                "Achieved",
                "Met",
                "You have effective organisational security management led at board level and articulated clearly in corresponding policies.",
            ],
        )
