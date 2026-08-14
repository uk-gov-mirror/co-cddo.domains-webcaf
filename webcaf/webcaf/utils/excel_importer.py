"""Converts a completed CAF Excel workbook back into assessment JSON.

Workbooks produced by :mod:`webcaf.webcaf.utils.excel_exporter` contain a
hidden mapping sheet (``JSON_MAP_SHEET_NAME``) whose rows record, for every
answer cell in the visible sheets, the JSON pointer path the value belongs to,
the type of value expected and whether the cell is required. The functions in
this module read that mapping, validate and transform the cell values, and
assemble the nested ``assessments_data`` dictionary stored on an Assessment.
"""

from typing import Any

from openpyxl import load_workbook

JSON_MAP_SHEET_NAME = "__webcaf_json_map"
JSON_MAP_HEADERS = ["visible_sheet", "visible_cell", "json_path", "value_type", "required"]
META_SHEET_NAME = "__webcaf_meta"

OUTCOME_STATUS_TO_REVIEW_DECISION = {
    "Achieved": "achieved",
    "Partially achieved": "partially-achieved",
    "Not achieved": "not-achieved",
}


class ExcelImportError(Exception):
    pass


def excel_to_assessment_json(excel_file) -> dict[str, Any]:
    """Convert an uploaded workbook to assessment JSON, strictly.

    Used when importing for real: raises :class:`ExcelImportError` if any cell
    holds an invalid value *or* if any required cell is blank.

    :param excel_file: A file path or file-like object readable by openpyxl.
    :return: The nested assessment data dictionary, keyed by outcome code.
    """
    data, _raw_rows, missing_required = _parse_excel(excel_file)
    if missing_required:
        raise ExcelImportError("; ".join(f"Required cell {cell} is blank" for cell in missing_required))
    return data


def excel_to_assessment_json_with_raw(excel_file) -> tuple[dict[str, Any], list[tuple[str, str, bool, Any]]]:
    """Convert an uploaded workbook to assessment JSON, leniently, for previewing.

    Unlike :func:`excel_to_assessment_json`, blank required cells do not raise:
    they are simply absent from the returned data, so the preview page can
    highlight them while still showing everything that was parsed.

    :param excel_file: A file path or file-like object readable by openpyxl.
    :return: A tuple of (assessment data dictionary, raw rows). Each raw row is
        a ``(json_path, value_type, required, raw_cell_value)`` tuple, one per
        mapped cell, in mapping-sheet order.
    """
    data, raw_rows, _missing_required = _parse_excel(excel_file)
    return data, raw_rows


def excel_framework_id(excel_file) -> Any:
    """
    Read the framework id the template was generated for.
    """
    wb = load_workbook(excel_file, data_only=True)
    if META_SHEET_NAME not in wb.sheetnames:
        return None
    for key, value in wb[META_SHEET_NAME].iter_rows(max_col=2, values_only=True):
        if key == "framework_id":
            return value
    return None


def excel_to_review_json(excel_file, framework: dict[str, Any]) -> dict[str, dict[str, dict[str, Any]]]:
    """
    Convert an uploaded workbook into review answers, grouped by objective and outcome.
    """
    assessment_data = excel_to_assessment_json(excel_file)

    outcome_to_objective = {
        outcome_code: objective["code"]
        for objective in framework.get("objectives", {}).values()
        for principle in objective.get("principles", {}).values()
        for outcome_code in principle.get("outcomes", {})
    }

    review_json: dict[str, dict[str, dict[str, Any]]] = {}
    errors: list[str] = []
    for outcome_code, outcome_data in assessment_data.items():
        objective_code = outcome_to_objective.get(outcome_code)
        if objective_code is None:
            errors.append(f"Outcome {outcome_code} is not part of the assessment's framework")
            continue

        outcome_answers: dict[str, Any] = {
            indicator_id: "yes" if agreed else "no"
            for indicator_id, agreed in outcome_data.get("indicators", {}).items()
        }
        confirmation = outcome_data.get("confirmation", {})
        outcome_answers["review_decision"] = OUTCOME_STATUS_TO_REVIEW_DECISION[confirmation["outcome_status"]]
        outcome_answers["review_comment"] = confirmation.get("confirm_outcome_confirm_comment", "")
        review_json.setdefault(objective_code, {})[outcome_code] = outcome_answers

    if errors:
        raise ExcelImportError("; ".join(errors))
    return review_json


def assessment_json_to_review_data(
    assessment_data: dict[str, Any],
    framework: dict[str, Any],
    stamped_by: str,
    stamped_by_role: str,
    stamped_by_email: str,
    stamped_at: str,
) -> dict[str, Any]:
    """
    Build the full review_data structure from parsed workbook data.
    Keys the workbook does not provide are created with empty values.
    """
    assessor_response_data: dict[str, Any] = {}
    for objective in framework.get("objectives", {}).values():
        objective_entry: dict[str, Any] = {}
        for principle in objective.get("principles", {}).values():
            for outcome_code, outcome in principle.get("outcomes", {}).items():
                outcome_data = assessment_data.get(outcome_code, {})
                answers = outcome_data.get("indicators", {})

                indicators: dict[str, Any] = {}
                for level, items in outcome.get("indicators", {}).items():
                    if not isinstance(items, dict):
                        continue
                    for item_code in items:
                        key = f"{level}_{item_code}"
                        answered = answers.get(key)
                        indicators[key] = "" if answered is None else ("yes" if answered else "no")
                        indicators[f"{key}_comment"] = ""

                confirmation = outcome_data.get("confirmation", {})
                objective_entry[outcome_code] = {
                    "indicators": indicators,
                    "review_data": {
                        "review_comment": confirmation.get("confirm_outcome_confirm_comment", ""),
                        "review_decision": OUTCOME_STATUS_TO_REVIEW_DECISION.get(
                            confirmation.get("outcome_status"), ""
                        ),
                    },
                    "recommendations": [],
                }
        objective_entry["recommendations"] = []
        objective_entry["objective-areas-of-improvement"] = ""
        objective_entry["objective-areas-of-good-practice"] = ""
        assessor_response_data[objective["code"]] = objective_entry

    assessor_response_data["system_and_scope"] = {
        "completed": "",
        "completed_data": {
            "review_details": {
                "caf_version": "",
                "review_type": "",
                "government_caf_profile": "",
                "self_assessment_reference_number": "",
            },
            "system_details": {
                "system_name": "",
                "system_ownership": "",
                "corporate_services": "",
                "system_description": "",
                "hosting_and_connectivity": "",
                "other_corporate_services": "",
                "previous_govassure_self_assessments": "",
            },
        },
    }
    assessor_response_data["additional_information"] = {
        "iar_period": {
            "start_date": "",
            "end_date": "",
        },
        "review_method": "",
        "quality_of_evidence": "",
        "company_details": {
            "company_name": "",
            "lead_assessor_name": "",
            "lead_assessor_email": "",
        },
        "areas_for_improvement": "",
        "areas_of_good_practice": "",
    }

    return {
        "review_finalised": {
            "review_finalised_at": stamped_at,
            "review_finalised_by": stamped_by,
            "review_finalised_by_role": stamped_by_role,
            "review_finalised_by_email": stamped_by_email,
        },
        "review_completion": {
            "review_completed": "yes",
            "review_completed_at": stamped_at,
            "review_completed_by": stamped_by,
            "review_completed_by_role": stamped_by_role,
            "review_completed_by_email": stamped_by_email,
        },
        "assessor_response_data": assessor_response_data,
    }


def _parse_excel(excel_file) -> tuple[dict[str, Any], list[tuple[str, str, bool, Any]], list[str]]:
    """Parse the Excel workbook into assessment data using its hidden mapping sheet.

    Walks every row of the ``JSON_MAP_SHEET_NAME`` sheet, reads the visible
    cell it points at, transforms the value according to the row's
    ``value_type`` (see :func:`_transform_value`) and writes it into the result
    dictionary at the row's JSON pointer path.

    Invalid values are collected and raised together as one
    :class:`ExcelImportError`. Blank required cells are *not* fatal here: they
    are returned in ``missing_required`` so callers can decide whether to raise
    (strict import) or flag them (preview).

    :param excel_file: A file path or file-like object readable by openpyxl.
    :return: A tuple of (assessment data dictionary, raw rows as described in
        :func:`excel_to_assessment_json_with_raw`, list of blank required cells
        as ``"Sheet!CELL"`` references).
    :raises ExcelImportError: If the workbook has no mapping sheet, the mapping
        sheet's header row is invalid, a mapped sheet is missing, or any cell
        value fails transformation.
    """
    wb = load_workbook(excel_file, data_only=True)
    if JSON_MAP_SHEET_NAME not in wb.sheetnames:
        raise ExcelImportError(f"Workbook is missing the {JSON_MAP_SHEET_NAME} sheet")

    map_ws = wb[JSON_MAP_SHEET_NAME]
    headers = [cell.value for cell in map_ws[1]]
    if headers != JSON_MAP_HEADERS:
        raise ExcelImportError(f"{JSON_MAP_SHEET_NAME} has an invalid header row")

    data: dict[str, Any] = {}
    raw_rows: list[tuple[str, str, bool, Any]] = []
    missing_required: list[str] = []
    errors: list[str] = []

    for visible_sheet, visible_cell, json_path, value_type, required in map_ws.iter_rows(min_row=2, values_only=True):
        if not visible_sheet or not visible_cell or not json_path:
            continue
        if visible_sheet not in wb.sheetnames:
            errors.append(f"Mapped sheet {visible_sheet} does not exist")
            continue

        value = wb[visible_sheet][visible_cell].value
        required_bool = bool(required)
        raw_rows.append((str(json_path), str(value_type), required_bool, value))

        # A blank required cell is not fatal here: record it so the preview can flag the row,
        # but skip transforming/storing it (there is nothing to store).
        if required_bool and _is_blank(value):
            missing_required.append(f"{visible_sheet}!{visible_cell}")
            continue

        try:
            transformed_value = _transform_value(value, str(value_type))
        except ExcelImportError as ex:
            errors.append(f"{visible_sheet}!{visible_cell}: {ex}")
            continue

        _set_json_pointer(data, str(json_path), transformed_value)

    if errors:
        raise ExcelImportError("; ".join(errors))

    _add_confirmation_defaults(data)
    return data, raw_rows, missing_required


def _transform_value(value: Any, value_type: str) -> Any:
    """Validate and convert a raw cell value according to its mapped ``value_type``.

    * ``indicator_answer`` — dropdown answers become booleans (blank counts as
      not agreed / ``False``).
    * ``outcome_status`` — must be one of the achievement statuses; returned as-is.
    * ``text`` — free text; blank becomes an empty string.

    :raises ExcelImportError: If the value is not valid for the type, or the
        type itself is unknown.
    """
    if value_type == "indicator_answer":
        if _is_blank(value):
            return False
        if isinstance(value, str):
            normalised = value.strip().lower()
            if normalised == "yes":
                return True
            if normalised == "no":
                return False
        raise ExcelImportError(f"invalid indicator answer {value!r}")

    if value_type == "outcome_status":
        if value in {"Achieved", "Partially achieved", "Not achieved"}:
            return value
        raise ExcelImportError(f"invalid outcome status {value!r}")

    if value_type == "text":
        return "" if _is_blank(value) else str(value)

    raise ExcelImportError(f"unknown value type {value_type!r}")


def _set_json_pointer(data: dict[str, Any], path: str, value: Any) -> None:
    """Write ``value`` into ``data`` at the JSON pointer ``path`` (e.g. ``/A1.a/indicators/x``),
    creating intermediate dictionaries as needed."""
    parts = [part for part in path.split("/") if part]
    if not parts:
        raise ExcelImportError("JSON path cannot be empty")

    current = data
    for part in parts[:-1]:
        current = current.setdefault(part, {})
    current[parts[-1]] = value


def _add_confirmation_defaults(data: dict[str, Any]) -> None:
    """Default each outcome's ``confirm_outcome`` to "confirm", as the form journey would,
    since the spreadsheet has no equivalent input."""
    for outcome_data in data.values():
        if isinstance(outcome_data, dict) and "confirmation" in outcome_data:
            outcome_data["confirmation"].setdefault("confirm_outcome", "confirm")


def _is_blank(value: Any) -> bool:
    return value is None or value == ""
